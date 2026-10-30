from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pathlib import Path


def render_pptx(content, output_path):

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    for slide_data in content["slides"]:

        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        # Title
        title_box = slide.shapes.add_textbox(
            Inches(0.7), Inches(0.4),
            Inches(12), Inches(0.8)
        )

        title_frame = title_box.text_frame
        title_frame.text = slide_data["title"]

        title_run = title_frame.paragraphs[0].runs[0]
        title_run.font.size = Pt(28)
        title_run.font.bold = True

        # Content
        content_box = slide.shapes.add_textbox(
            Inches(0.9), Inches(1.7),
            Inches(11.5), Inches(4.8)
        )

        text_frame = content_box.text_frame

        for item in slide_data.get("content", []):

            paragraph = text_frame.add_paragraph()
            paragraph.text = item
            paragraph.font.size = Pt(20)
            paragraph.space_after = Pt(12)

    Path(output_path).parent.mkdir(
        parents=True, exist_ok=True
    )

    prs.save(output_path)

    return output_path