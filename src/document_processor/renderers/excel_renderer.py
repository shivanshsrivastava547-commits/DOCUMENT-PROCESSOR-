from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from pathlib import Path


def render_excel(content, output_path):

    wb = Workbook()

    for index, sheet_data in enumerate(content["sheets"]):

        if index == 0:
            ws = wb.active
            ws.title = sheet_data["name"]
        else:
            ws = wb.create_sheet(sheet_data["name"])

        headers = sheet_data["headers"]
        rows = sheet_data["rows"]

        ws.append(headers)

        for row in rows:
            ws.append(row)

        # Header formatting
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                "solid", fgColor="1F4E78"
            )
            cell.alignment = Alignment(horizontal="center")

        # Column widths
        for column in ws.columns:
            max_length = max(
                len(str(cell.value or ""))
                for cell in column
            )
            ws.column_dimensions[
                column[0].column_letter
            ].width = min(max_length + 4, 35)

    Path(output_path).parent.mkdir(
        parents=True, exist_ok=True
    )

    wb.save(output_path)

    return output_path